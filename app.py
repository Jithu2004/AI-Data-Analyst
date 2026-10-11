import json
import os
import re
from pathlib import Path
from datetime import datetime

import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

from dotenv import load_dotenv
from google import genai
from pydantic import BaseModel, Field

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# =========================================================
# 1. CONFIGURATION
# =========================================================

load_dotenv()

API_KEY = os.getenv("GEMINI_API_KEY")

if not API_KEY:
    st.error("GEMINI_API_KEY not found in .env file.")
    st.stop()

client = genai.Client(api_key=API_KEY)

MODEL_NAME = "gemini-3.5-flash-lite"


# =========================================================
# 2. PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="AI Data Analyst",
    page_icon="📊",
    layout="wide"
)

st.title("📊 AI Data Analyst")

st.caption(
    "Upload one or more compatible datasets and ask questions "
    "using natural language."
)


# =========================================================
# 3. SESSION MEMORY
# =========================================================

if "last_question" not in st.session_state:
    st.session_state.last_question = ""

if "last_plan" not in st.session_state:
    st.session_state.last_plan = None

if "last_result" not in st.session_state:
    st.session_state.last_result = None

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

# Separate history for report Q&A.
if "report_history" not in st.session_state:
    st.session_state.report_history = []


# =========================================================
# 4. ANALYSIS MODELS
# =========================================================

class FilterCondition(BaseModel):
    column: str
    value: str


class AnalysisPlan(BaseModel):
    operation: str
    metric: str | None = None
    filters: list[FilterCondition] = Field(default_factory=list)
    group_by: str | None = None
    sort_direction: str | None = None
    top_n: int | None = None


# =========================================================
# 5. TEXT HELPERS
# =========================================================

def normalize_text(value):
    return str(value).strip().lower()


def clean_json_response(text):

    text = text.strip()

    if text.startswith("```"):

        lines = text.splitlines()

        if lines:
            lines = lines[1:]

        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]

        text = "\n".join(lines)

    return text.strip()


# =========================================================
# 6. NORMALIZE GEMINI PLAN
# =========================================================

def normalize_gemini_plan(data):

    if not isinstance(data, dict):
        data = {}

    operation = data.get(
        "operation",
        "sum"
    )

    if isinstance(operation, list):
        operation = (
            operation[0]
            if operation
            else "sum"
        )

    operation = str(
        operation
    ).strip().lower()

    metric = data.get(
        "metric",
        "profit"
    )

    if isinstance(metric, list):
        metric = (
            metric[0]
            if metric
            else "profit"
        )

    if metric is not None:
        metric = str(
            metric
        ).strip().lower()

    group_by = data.get(
        "group_by"
    )

    if isinstance(group_by, list):
        group_by = (
            group_by[0]
            if group_by
            else None
        )

    if group_by is not None:
        group_by = str(
            group_by
        ).strip()

    sort_direction = data.get(
        "sort_direction"
    )

    if isinstance(sort_direction, list):
        sort_direction = (
            sort_direction[0]
            if sort_direction
            else None
        )

    if sort_direction is not None:
        sort_direction = str(
            sort_direction
        ).strip().lower()

    top_n = data.get(
        "top_n"
    )

    if isinstance(top_n, list):
        top_n = (
            top_n[0]
            if top_n
            else None
        )

    if top_n is not None:

        try:
            top_n = int(top_n)

        except Exception:
            top_n = None

    raw_filters = data.get(
        "filters",
        []
    )

    if not isinstance(
        raw_filters,
        list
    ):
        raw_filters = []

    cleaned_filters = []

    for item in raw_filters:

        if not isinstance(
            item,
            dict
        ):
            continue

        column = item.get(
            "column"
        )

        value = item.get(
            "value"
        )

        if isinstance(
            column,
            list
        ):
            column = (
                column[0]
                if column
                else None
            )

        if isinstance(
            value,
            list
        ):
            value = (
                value[0]
                if value
                else None
            )

        if column is None:
            continue

        if value is None:
            continue

        cleaned_filters.append(
            {
                "column": str(
                    column
                ).strip(),

                "value": str(
                    value
                ).strip()
            }
        )

    return {
        "operation": operation,
        "metric": metric,
        "filters": cleaned_filters,
        "group_by": group_by,
        "sort_direction": sort_direction,
        "top_n": top_n
    }


# =========================================================
# 7. READ UPLOADED FILE
# =========================================================

def read_uploaded_file(
    uploaded_file
):

    filename = uploaded_file.name.lower()

    try:

        if filename.endswith(".csv"):

            return pd.read_csv(
                uploaded_file
            )

        if filename.endswith(".xlsx"):

            return pd.read_excel(
                uploaded_file
            )

        return None

    except Exception as e:

        st.error(
            f"Could not read "
            f"{uploaded_file.name}: {e}"
        )

        return None


# =========================================================
# 8. NORMALIZE DATAFRAME
# =========================================================

def normalize_dataframe(df):

    df = df.copy()

    df = df.dropna(
        axis=1,
        how="all"
    )

    df.columns = [
        str(column).strip()
        for column in df.columns
    ]

    numeric_columns = [
        "sales",
        "revenue",
        "profit",
        "loss",
        "unsold"
    ]

    for column in numeric_columns:

        if column in df.columns:

            df[column] = pd.to_numeric(
                df[column],
                errors="coerce"
            ).fillna(0)

    return df


# =========================================================
# 9. CHECK SCHEMA
# =========================================================

def check_schema_compatibility(
    dataframes
):

    if not dataframes:
        return True, ""

    reference_columns = list(
        dataframes[0].columns
    )

    reference_set = set(
        reference_columns
    )

    for index, df in enumerate(
        dataframes[1:],
        start=2
    ):

        current_set = set(
            df.columns
        )

        missing = (
            reference_set
            - current_set
        )

        extra = (
            current_set
            - reference_set
        )

        if missing or extra:

            message = (
                f"File {index} has "
                f"an incompatible schema.\n\n"
            )

            if missing:

                message += (
                    "Missing columns: "
                    + ", ".join(
                        sorted(missing)
                    )
                    + "\n"
                )

            if extra:

                message += (
                    "Additional columns: "
                    + ", ".join(
                        sorted(extra)
                    )
                )

            return False, message

    return True, ""


# =========================================================
# 10. DATASET CONTEXT
# =========================================================

def create_dataset_context(df):

    context = {
        "columns": list(
            df.columns
        ),

        "dtypes": {
            column: str(dtype)
            for column, dtype
            in df.dtypes.items()
        },

        "categorical_values": {}
    }

    categorical_columns = [
        "year_month",
        "state",
        "city",
        "product",
        "brand",
        "product_series"
    ]

    for column in categorical_columns:

        if column in df.columns:

            values = (
                df[column]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            context[
                "categorical_values"
            ][column] = values

    return context


# =========================================================
# 11. FIND DATASET VALUE
# =========================================================

def find_dataset_value(
    question,
    df,
    column
):

    if column not in df.columns:
        return None

    question_normalized = (
        normalize_text(
            question
        )
    )

    values = (
        df[column]
        .dropna()
        .astype(str)
        .unique()
        .tolist()
    )

    values.sort(
        key=lambda value:
        len(
            normalize_text(
                value
            )
        ),
        reverse=True
    )

    for value in values:

        normalized_value = (
            normalize_text(
                value
            )
        )

        if normalized_value in question_normalized:

            return value

    return None


# =========================================================
# 12. DETECT EXPLICIT FILTERS
# =========================================================

def detect_explicit_filters(
    question,
    df
):

    filters = []

    columns = [
        "year_month",
        "state",
        "city",
        "product",
        "brand",
        "product_series"
    ]

    for column in columns:

        value = find_dataset_value(
            question,
            df,
            column
        )

        if value is not None:

            filters.append(
                FilterCondition(
                    column=column,
                    value=value
                )
            )

    return filters


# =========================================================
# 13. TOP N DETECTION
# =========================================================

def detect_top_n(question):

    question_normalized = (
        normalize_text(
            question
        )
    )

    patterns = [
        r"\btop\s+(\d+)\b",
        r"\bfirst\s+(\d+)\b",
        r"\bbest\s+(\d+)\b"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            question_normalized
        )

        if match:
            return int(
                match.group(1)
            )

    if (
        "most sold" in question_normalized
        or "most sales" in question_normalized
        or "highest sold" in question_normalized
        or "highest sales" in question_normalized
        or "best selling" in question_normalized
        or "highest selling" in question_normalized
    ):
        return 1

    return None


# =========================================================
# 14. BOTTOM N DETECTION
# =========================================================

def detect_bottom_n(question):

    question_normalized = (
        normalize_text(
            question
        )
    )

    patterns = [
        r"\bbottom\s+(\d+)\b",
        r"\blast\s+(\d+)\b",
        r"\bworst\s+(\d+)\b"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            question_normalized
        )

        if match:
            return int(
                match.group(1)
            )

    if (
        "least sold" in question_normalized
        or "least sales" in question_normalized
        or "lowest sold" in question_normalized
        or "lowest sales" in question_normalized
        or "least selling" in question_normalized
        or "lowest selling" in question_normalized
    ):
        return 1

    return None


# =========================================================
# 15. IMPROVE ANALYSIS PLAN
# =========================================================

def improve_analysis_plan(
    plan,
    question,
    df
):

    valid_metrics = {
        "sales",
        "revenue",
        "profit",
        "loss",
        "unsold",
        "profit_margin",
        "revenue_per_sale",
        "profit_per_sale"
    }

    metric = (
        normalize_text(
            plan.metric
        )
        if plan.metric
        else None
    )

    if metric not in valid_metrics:
        metric = "profit"

    valid_operations = {
        "sum",
        "mean",
        "count",
        "max",
        "min"
    }

    operation = normalize_text(
        plan.operation
    )

    if operation not in valid_operations:
        operation = "sum"

    sort_direction = (
        normalize_text(
            plan.sort_direction
        )
        if plan.sort_direction
        else None
    )

    if sort_direction not in {
        "asc",
        "desc"
    }:
        sort_direction = None

    question_normalized = normalize_text(
        question
    )

    group_patterns = {
        "state": [
            "by state",
            "per state",
            "state wise",
            "state-wise",
            "statewise",
            "states"
        ],

        "brand": [
            "by brand",
            "per brand",
            "brand wise",
            "brand-wise",
            "brandwise",
            "brands"
        ],

        "product": [
            "by product",
            "per product",
            "product wise",
            "product-wise",
            "productwise",
            "products"
        ],

        "city": [
            "by city",
            "per city",
            "city wise",
            "city-wise",
            "citywise",
            "cities"
        ],

        "product_series": [
            "by product series",
            "per product series",
            "product series wise",
            "product-series wise"
        ],

        "year_month": [
            "by month",
            "per month",
            "monthly",
            "month wise",
            "month-wise",
            "by year month",
            "by year_month"
        ]
    }

    explicit_group_by = None

    for column, patterns in group_patterns.items():

        if any(
            pattern in question_normalized
            for pattern in patterns
        ):

            explicit_group_by = column
            break

    if explicit_group_by is not None:

        group_by = explicit_group_by

    else:

        group_by = plan.group_by

        if group_by:

            group_by = group_by.strip()

            if group_by not in df.columns:
                group_by = None

    # -----------------------------------------------------
    # IMPORTANT:
    # If the current question explicitly asks for a new
    # grouping dimension, do not carry the previous
    # question's unrelated filters into the new analysis.
    # -----------------------------------------------------

    if explicit_group_by is not None:

        final_filters = []

    else:

        final_filters = []

        for condition in plan.filters:

            if condition.column not in df.columns:
                continue

            final_filters.append(
                FilterCondition(
                    column=condition.column,
                    value=str(
                        condition.value
                    )
                )
            )

    # =====================================================
    # METRIC DETECTION
    # =====================================================

    metric_keywords = [
        (
            "profit margin",
            "profit_margin"
        ),

        (
            "profit per sale",
            "profit_per_sale"
        ),

        (
            "revenue per sale",
            "revenue_per_sale"
        ),

        (
            "revenue",
            "revenue"
        ),

        (
            "profit",
            "profit"
        ),

        (
            "sales",
            "sales"
        ),

        (
            "loss",
            "loss"
        ),

        (
            "unsold",
            "unsold"
        )
    ]

    for keyword, detected_metric in (
        metric_keywords
    ):

        if keyword in question_normalized:

            metric = detected_metric
            break

    # =====================================================
    # OPERATION DETECTION
    # =====================================================

    if (
        "average" in question_normalized
        or "avg" in question_normalized
        or "mean" in question_normalized
    ):

        operation = "mean"

    elif (
        "maximum value"
        in question_normalized
        or
        "max value"
        in question_normalized
    ):

        operation = "max"

    elif (
        "minimum value"
        in question_normalized
        or
        "min value"
        in question_normalized
    ):

        operation = "min"

    elif (
        "count"
        in question_normalized
        or
        "number of"
        in question_normalized
        or
        "how many"
        in question_normalized
    ):

        operation = "count"

    elif (
        "total"
        in question_normalized
        or
        "sum"
        in question_normalized
    ):

        operation = "sum"

    # =====================================================
    # TOP / BOTTOM N
    # =====================================================

    top_n = plan.top_n

    detected_top_n = (
        detect_top_n(
            question
        )
    )

    if detected_top_n is not None:

        top_n = detected_top_n
        sort_direction = "desc"

    detected_bottom_n = (
        detect_bottom_n(
            question
        )
    )

    if detected_bottom_n is not None:

        top_n = detected_bottom_n
        sort_direction = "asc"

    # =====================================================
    # STATE RANKING
    # =====================================================

    if (
        "least state"
        in question_normalized
        or
        "lowest state"
        in question_normalized
        or
        "lowest state by"
        in question_normalized
    ):

        group_by = "state"
        top_n = 1
        sort_direction = "asc"

        # This is a comparison across ALL states.
        # Do not keep an old state filter.
        final_filters = []

    elif (
        "most state"
        in question_normalized
        or
        "highest state"
        in question_normalized
        or
        "highest state by"
        in question_normalized
    ):

        group_by = "state"
        top_n = 1
        sort_direction = "desc"

        # This is a comparison across ALL states.
        # Do not keep an old state filter.
        final_filters = []

    # =====================================================
    # PRODUCT RANKING
    # =====================================================

    if (
        "most sold product"
        in question_normalized
        or
        "most selling product"
        in question_normalized
        or
        "best selling product"
        in question_normalized
        or
        "highest selling product"
        in question_normalized
        or
        "product sold the most"
        in question_normalized
    ):

        metric = "sales"
        group_by = "product"
        top_n = 1
        sort_direction = "desc"

        final_filters = []

    elif (
        "least sold product"
        in question_normalized
        or
        "least selling product"
        in question_normalized
        or
        "lowest selling product"
        in question_normalized
        or
        "lowest sales product"
        in question_normalized
        or
        "product sold the least"
        in question_normalized
    ):

        metric = "sales"
        group_by = "product"
        top_n = 1
        sort_direction = "asc"

        final_filters = []

    # =====================================================
    # EXPLICIT FILTERS
    # =====================================================

    detected_filters = (
        detect_explicit_filters(
            question,
            df
        )
    )

    for detected in detected_filters:

        replaced = False

        for index, existing in enumerate(
            final_filters
        ):

            if (
                existing.column
                == detected.column
            ):

                final_filters[index] = (
                    detected
                )

                replaced = True
                break

        if not replaced:

            final_filters.append(
                detected
            )

    # =====================================================
    # TOP N VALIDATION
    # =====================================================

    if top_n is not None:

        try:

            top_n = int(
                top_n
            )

            if top_n <= 0:
                top_n = None

        except Exception:

            top_n = None

    return AnalysisPlan(
        operation=operation,
        metric=metric,
        filters=final_filters,
        group_by=group_by,
        sort_direction=sort_direction,
        top_n=top_n
    )


# =========================================================
# 16. FOLLOW-UP CONTEXT
# =========================================================

def resolve_follow_up(
    question,
    previous_plan,
    previous_question,
    df
):

    if previous_plan is None:
        return None

    question_normalized = normalize_text(
        question
    )

    # =====================================================
    # IMPORTANT FIX
    #
    # Do NOT treat every new question as a follow-up.
    #
    # "profit by brand" after
    # "least state in revenue"
    #
    # is a NEW question.
    #
    # But:
    # "top 3"
    # "what about ASUS"
    # "what about Karnataka"
    #
    # can be genuine follow-ups.
    # =====================================================

    follow_up_phrases = [
        "what about",
        "how about",
        "compare them",
        "compare these",
        "compare it",
        "and what about",
        "what about it",
        "what about them",
        "same for",
        "same analysis for",
        "now for",
        "instead",
        "also show",
        "also give me"
    ]

    is_follow_up = any(
        phrase in question_normalized
        for phrase in follow_up_phrases
    )

    top_n = detect_top_n(
        question
    )

    bottom_n = detect_bottom_n(
        question
    )

    # A completely new question must use Gemini's fresh
    # analysis plan.
    if (
        top_n is None
        and
        bottom_n is None
        and
        not is_follow_up
    ):

        return None

    plan = AnalysisPlan(
        **previous_plan
    )

    # =====================================================
    # TOP N FOLLOW-UP
    # =====================================================

    if top_n is not None:

        plan.top_n = top_n
        plan.sort_direction = "desc"

    # =====================================================
    # BOTTOM N FOLLOW-UP
    # =====================================================

    if bottom_n is not None:

        plan.top_n = bottom_n
        plan.sort_direction = "asc"

    # =====================================================
    # METRIC CHANGES
    # =====================================================

    metric_keywords = [
        (
            "profit margin",
            "profit_margin"
        ),

        (
            "profit per sale",
            "profit_per_sale"
        ),

        (
            "revenue per sale",
            "revenue_per_sale"
        ),

        (
            "revenue",
            "revenue"
        ),

        (
            "profit",
            "profit"
        ),

        (
            "sales",
            "sales"
        ),

        (
            "loss",
            "loss"
        ),

        (
            "unsold",
            "unsold"
        )
    ]

    for keyword, metric in (
        metric_keywords
    ):

        if keyword in question_normalized:

            plan.metric = metric
            break

    # =====================================================
    # EXPLICIT GROUPING
    # =====================================================

    group_patterns = {
        "state": [
            "by state",
            "per state",
            "state wise",
            "state-wise",
            "statewise",
            "states"
        ],

        "brand": [
            "by brand",
            "per brand",
            "brand wise",
            "brand-wise",
            "brandwise",
            "brands"
        ],

        "product": [
            "by product",
            "per product",
            "product wise",
            "product-wise",
            "productwise",
            "products"
        ],

        "city": [
            "by city",
            "per city",
            "city wise",
            "city-wise",
            "citywise",
            "cities"
        ],

        "product_series": [
            "by product series",
            "per product series",
            "product series wise",
            "product-series wise"
        ],

        "year_month": [
            "by month",
            "per month",
            "monthly",
            "month wise",
            "month-wise",
            "by year month",
            "by year_month"
        ]
    }

    explicit_group_by = None

    for column, patterns in (
        group_patterns.items()
    ):

        if any(
            pattern in question_normalized
            for pattern in patterns
        ):

            explicit_group_by = column
            break

    if explicit_group_by is not None:

        plan.group_by = (
            explicit_group_by
        )

        # New explicit grouping = new dimension.
        # Remove old unrelated filters.
        plan.filters = []

        if (
            top_n is None
            and
            bottom_n is None
        ):

            plan.top_n = None
            plan.sort_direction = None

    # =====================================================
    # EXPLICIT FILTERS
    # =====================================================

    detected_filters = (
        detect_explicit_filters(
            question,
            df
        )
    )

    for detected in detected_filters:

        replaced = False

        for index, existing in enumerate(
            plan.filters
        ):

            if (
                existing.column
                == detected.column
            ):

                plan.filters[index] = (
                    detected
                )

                replaced = True
                break

        if not replaced:

            plan.filters.append(
                detected
            )

    # =====================================================
    # PRODUCT RANKING FOLLOW-UP
    # =====================================================

    if (
        "most sold product"
        in question_normalized
        or
        "most selling product"
        in question_normalized
        or
        "best selling product"
        in question_normalized
        or
        "highest selling product"
        in question_normalized
        or
        "product sold the most"
        in question_normalized
    ):

        plan.metric = "sales"
        plan.group_by = "product"
        plan.top_n = 1
        plan.sort_direction = "desc"
        plan.filters = []

    elif (
        "least sold product"
        in question_normalized
        or
        "least selling product"
        in question_normalized
        or
        "lowest selling product"
        in question_normalized
        or
        "lowest sales product"
        in question_normalized
        or
        "product sold the least"
        in question_normalized
    ):

        plan.metric = "sales"
        plan.group_by = "product"
        plan.top_n = 1
        plan.sort_direction = "asc"
        plan.filters = []

    return plan


# =========================================================
# 17. APPLY FILTERS
# =========================================================

def apply_filters(
    df,
    filters
):

    filtered_df = df.copy()

    for condition in filters:

        column = condition.column

        value = normalize_text(
            condition.value
        )

        if column not in filtered_df.columns:
            continue

        mask = (
            filtered_df[column]
            .astype(str)
            .str.strip()
            .str.lower()
            == value
        )

        filtered_df = filtered_df[
            mask
        ]

    return filtered_df


# =========================================================
# 18. CALCULATE METRIC
# =========================================================

def calculate_metric_series(
    df,
    metric
):

    if metric in {
        "sales",
        "revenue",
        "profit",
        "loss",
        "unsold"
    }:

        return pd.to_numeric(
            df[metric],
            errors="coerce"
        ).fillna(0)

    if metric == "profit_margin":

        revenue = pd.to_numeric(
            df["revenue"],
            errors="coerce"
        ).fillna(0)

        profit = pd.to_numeric(
            df["profit"],
            errors="coerce"
        ).fillna(0)

        return (
            profit
            .div(
                revenue.replace(
                    0,
                    pd.NA
                )
            )
            .fillna(0)
            * 100
        )

    if metric == "revenue_per_sale":

        revenue = pd.to_numeric(
            df["revenue"],
            errors="coerce"
        ).fillna(0)

        sales = pd.to_numeric(
            df["sales"],
            errors="coerce"
        ).fillna(0)

        return (
            revenue
            .div(
                sales.replace(
                    0,
                    pd.NA
                )
            )
            .fillna(0)
        )

    if metric == "profit_per_sale":

        profit = pd.to_numeric(
            df["profit"],
            errors="coerce"
        ).fillna(0)

        sales = pd.to_numeric(
            df["sales"],
            errors="coerce"
        ).fillna(0)

        return (
            profit
            .div(
                sales.replace(
                    0,
                    pd.NA
                )
            )
            .fillna(0)
        )

    return pd.Series(
        [0] * len(df),
        index=df.index
    )


# =========================================================
# 19. RUN ANALYSIS
# =========================================================

def run_analysis_plan(
    plan,
    df
):

    filtered_df = apply_filters(
        df,
        plan.filters
    )

    matched_rows = len(
        filtered_df
    )

    metric = (
        plan.metric
        or "profit"
    )

    metric_values = (
        calculate_metric_series(
            filtered_df,
            metric
        )
    )

    if plan.group_by:

        working_df = (
            filtered_df.copy()
        )

        working_df[
            "_metric_value"
        ] = metric_values

        grouped = (
            working_df
            .groupby(
                plan.group_by
            )["_metric_value"]
        )

        if plan.operation == "sum":
            result = grouped.sum()

        elif plan.operation == "mean":
            result = grouped.mean()

        elif plan.operation == "count":
            result = grouped.count()

        elif plan.operation == "max":
            result = grouped.max()

        elif plan.operation == "min":
            result = grouped.min()

        else:
            result = grouped.sum()

        result = result.reset_index()

        result.columns = [
            plan.group_by,
            metric
        ]

        if plan.sort_direction == "asc":

            result = result.sort_values(
                by=metric,
                ascending=True
            )

        else:

            result = result.sort_values(
                by=metric,
                ascending=False
            )

        if plan.top_n:

            result = result.head(
                plan.top_n
            )

        result = result.reset_index(
            drop=True
        )

        return result, matched_rows

    if plan.operation == "sum":
        value = metric_values.sum()

    elif plan.operation == "mean":
        value = metric_values.mean()

    elif plan.operation == "count":
        value = metric_values.count()

    elif plan.operation == "max":
        value = metric_values.max()

    elif plan.operation == "min":
        value = metric_values.min()

    else:
        value = metric_values.sum()

    result = pd.DataFrame(
        {
            metric: [value]
        }
    )

    return result, matched_rows


# =========================================================
# 20. CREATE CHART
# =========================================================

def create_chart(result):

    if len(result.columns) < 2:
        return

    if len(result) == 0:
        return

    x_column = result.columns[0]
    y_column = result.columns[1]

    chart_data = result.copy()

    chart_data = chart_data.sort_values(
        by=y_column,
        ascending=False
    )

    if len(chart_data) > 20:
        chart_data = chart_data.head(20)

    chart_data = chart_data.set_index(
        x_column
    )[[y_column]]

    st.bar_chart(
        chart_data,
        horizontal=True
    )


# =========================================================
# 21. BUSINESS INSIGHTS
# =========================================================

def generate_business_insights(
    question,
    result
):

    result_text = result.to_string(
        index=False
    )

    prompt = (
        "You are a business analyst.\n\n"

        "User question:\n"
        + str(question)

        + "\n\n"

        "Actual calculated result:\n"
        + result_text

        + "\n\n"

        "Generate 3 to 5 concise business insights.\n\n"

        "Rules:\n"
        "- Use only the provided result.\n"
        "- Do not invent numbers.\n"
        "- Do not perform new calculations.\n"
        "- Do not claim facts that are not visible.\n"
        "- Keep the language simple.\n"
        "- Return bullet points only."
    )

    try:

        response = (
            client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt
            )
        )

        return response.text.strip()

    except Exception as e:

        return (
            "Could not generate insights: "
            + str(e)
        )


# =========================================================
# 22. ANOMALY DETECTION
# =========================================================

def detect_anomalies(result):

    if len(result) < 4:
        return None

    if len(result.columns) < 2:
        return None

    value_column = result.columns[1]

    values = pd.to_numeric(
        result[value_column],
        errors="coerce"
    )

    if values.isna().all():
        return None

    q1 = values.quantile(
        0.25
    )

    q3 = values.quantile(
        0.75
    )

    iqr = q3 - q1

    lower_bound = (
        q1 - 1.5 * iqr
    )

    upper_bound = (
        q3 + 1.5 * iqr
    )

    anomalies = result[
        (values < lower_bound)
        |
        (values > upper_bound)
    ].copy()

    if len(anomalies) == 0:
        return None

    anomalies[
        "anomaly_type"
    ] = anomalies[
        value_column
    ].apply(
        lambda x:
        "High"
        if x > upper_bound
        else "Low"
    )

    anomalies[
        "lower_bound"
    ] = lower_bound

    anomalies[
        "upper_bound"
    ] = upper_bound

    return anomalies


# =========================================================
# 23. EXPLAIN ANOMALIES
# =========================================================

def explain_anomalies(
    question,
    anomalies
):

    anomaly_text = (
        anomalies.to_string(
            index=False
        )
    )

    prompt = (
        "You are a business analyst.\n\n"

        "User question:\n"
        + str(question)

        + "\n\n"

        "Detected statistical anomalies:\n"
        + anomaly_text

        + "\n\n"

        "Explain these anomalies carefully.\n\n"

        "Rules:\n"
        "- Do not invent causes.\n"
        "- Say anomalies may require investigation.\n"
        "- Use only the information provided.\n"
        "- Keep the explanation concise.\n"
        "- Return 2 to 4 bullet points."
    )

    try:

        response = (
            client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt
            )
        )

        return response.text.strip()

    except Exception as e:

        return (
            "Could not explain anomalies: "
            + str(e)
        )


# =========================================================
# 24. GEMINI PLANNER
# =========================================================

def ask_gemini(
    question,
    df,
    previous_question="",
    previous_plan=None,
    previous_result=None
):

    dataset_context = (
        create_dataset_context(
            df
        )
    )

    previous_context = (
        "No previous analysis."
    )

    if previous_question:

        previous_context = (
            "Previous question:\n"
            + str(previous_question)

            + "\n\n"

            "Previous plan:\n"
            + json.dumps(
                previous_plan,
                indent=2
            )

            + "\n\n"

            "Previous result:\n"
            + str(previous_result)
        )

    prompt = (
        "You are the planning component "
        "of an AI Data Analyst.\n\n"

        "Convert the user's question into "
        "a structured analysis plan.\n\n"

        "Do NOT calculate numbers.\n\n"

        "Python performs the actual calculation.\n\n"

        "Dataset:\n"
        + json.dumps(
            dataset_context,
            indent=2
        )

        + "\n\n"

        "Previous conversation:\n"
        + previous_context

        + "\n\n"

        "Current question:\n"
        + str(question)

        + "\n\n"

        "Rules:\n"
        "1. Return ONLY valid JSON.\n"
        "2. No markdown.\n"
        "3. Do not calculate numerical answers.\n"
        "4. Use previous context for follow-up questions.\n"
        "5. Preserve previous metric unless a new metric is requested.\n"
        "6. Preserve previous group_by for top/bottom questions.\n"
        "7. top N = descending + top_n=N.\n"
        "8. bottom N = ascending + top_n=N.\n"
        "9. Most sold product = sales grouped by product, descending, top_n=1.\n"
        "10. Least sold product = sales grouped by product, ascending, top_n=1.\n"
        "11. what about ASUS = brand filter ASUS while preserving analysis.\n"
        "12. group_by MUST be a string or null.\n"
        "13. filters MUST be an array.\n"
        "14. top_n MUST be a number or null.\n"
        "15. Use only actual dataset columns.\n\n"

        "Allowed operations:\n"
        "sum\n"
        "mean\n"
        "count\n"
        "max\n"
        "min\n\n"

        "Allowed metrics:\n"
        "sales\n"
        "revenue\n"
        "profit\n"
        "loss\n"
        "unsold\n"
        "profit_margin\n"
        "revenue_per_sale\n"
        "profit_per_sale\n\n"

        "Return JSON in this format:\n"

        "{\n"
        '  "operation": "sum",\n'
        '  "metric": "profit",\n'
        '  "filters": [],\n'
        '  "group_by": "brand",\n'
        '  "sort_direction": "desc",\n'
        '  "top_n": 3\n'
        "}"
    )

    try:

        response = (
            client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt
            )
        )

        cleaned = clean_json_response(
            response.text
        )

        raw_data = json.loads(
            cleaned
        )

        data = normalize_gemini_plan(
            raw_data
        )

        return AnalysisPlan(
            **data
        )

    except Exception as e:

        st.error(
            "Gemini planning error: "
            + str(e)
        )

        return None


# =========================================================
# 25. WORD REPORT HELPERS
# =========================================================

def format_value(
    value,
    column_name=""
):

    try:

        numeric_value = float(
            value
        )

    except Exception:

        return str(value)

    column_name = (
        str(column_name)
        .lower()
    )

    if column_name in {
        "revenue",
        "profit",
        "loss",
        "revenue_per_sale",
        "profit_per_sale"
    }:

        return (
            "₹"
            + f"{numeric_value:,.2f}"
        )

    if column_name == "profit_margin":

        return (
            f"{numeric_value:,.2f}%"
        )

    if column_name in {
        "sales",
        "unsold"
    }:

        return (
            f"{numeric_value:,.0f}"
        )

    return (
        f"{numeric_value:,.2f}"
    )


def add_dataframe_to_word(
    document,
    dataframe,
    title=None
):

    if title:

        document.add_heading(
            title,
            level=2
        )

    if dataframe is None:
        return

    if len(dataframe) == 0:

        document.add_paragraph(
            "No data available."
        )

        return

    table = document.add_table(
        rows=1,
        cols=len(
            dataframe.columns
        )
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    header_cells = (
        table.rows[0].cells
    )

    for index, column in enumerate(
        dataframe.columns
    ):

        header_cells[index].text = (
            str(column)
            .replace(
                "_",
                " "
            )
            .title()
        )

    for _, row in (
        dataframe.iterrows()
    ):

        cells = (
            table.add_row().cells
        )

        for index, column in enumerate(
            dataframe.columns
        ):

            cells[index].text = (
                format_value(
                    row[column],
                    column
                )
            )


# =========================================================
# 26. REPORT HISTORY
# =========================================================

def get_report_history_for_output(
    include_previous_qa
):

    history = (
        st.session_state.report_history
    )

    if not history:
        return []

    # CHECKED:
    # Include every previous Q&A.
    if include_previous_qa:

        return list(history)

    # UNCHECKED:
    # Include ONLY the latest Q&A.
    return [
        history[-1]
    ]


# =========================================================
# 27. GENERATE ANALYST REPORT
# =========================================================

def generate_analyst_report(
    question,
    plan,
    result,
    matched_rows,
    insights,
    anomalies,
    anomaly_explanation
):

    result_text = (
        result.to_string(
            index=False
        )
    )

    if anomalies is not None:

        anomaly_text = (
            anomalies.to_string(
                index=False
            )
        )

    else:

        anomaly_text = "None"

    prompt = (
        "You are a senior business analyst.\n\n"

        "Create a concise analyst report using "
        "ONLY the information provided below.\n\n"

        "User question:\n"
        + str(question)

        + "\n\n"

        "Analysis plan:\n"
        + json.dumps(
            plan.model_dump(),
            indent=2
        )

        + "\n\n"

        "Rows matched:\n"
        + str(matched_rows)

        + "\n\n"

        "Actual calculated result:\n"
        + result_text

        + "\n\n"

        "AI business insights:\n"
        + str(insights)

        + "\n\n"

        "Statistical anomalies:\n"
        + anomaly_text

        + "\n\n"

        "Anomaly explanation:\n"
        + str(
            anomaly_explanation
        )

        + "\n\n"

        "Return ONLY valid JSON.\n\n"

        "Required structure:\n"

        "{\n"
        '  "executive_summary": "short summary",\n'
        '  "key_findings": ["finding 1", "finding 2"],\n'
        '  "top_performers": ["item 1", "item 2"],\n'
        '  "anomalies": ["item 1"],\n'
        '  "recommendations": ["recommendation 1", "recommendation 2"]\n'
        "}\n\n"

        "Rules:\n"
        "- Do not invent numbers.\n"
        "- Do not calculate new numbers.\n"
        "- Use only supplied information.\n"
        "- Keep findings concise.\n"
        "- Recommendations must be based on the supplied analysis."
    )

    try:

        response = (
            client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt
            )
        )

        cleaned = clean_json_response(
            response.text
        )

        report = json.loads(
            cleaned
        )

        if not isinstance(
            report,
            dict
        ):

            raise ValueError(
                "Invalid report format."
            )

        return report

    except Exception as e:

        return {
            "executive_summary": (
                "The analyst report could not "
                "be generated automatically."
            ),

            "key_findings": [
                "Analysis was completed successfully."
            ],

            "top_performers": [],

            "anomalies": [],

            "recommendations": [
                "Review the calculated analysis result."
            ],

            "error": str(e)
        }


# =========================================================
# 28. CREATE WORD REPORT
# =========================================================

def create_word_report(
    question,
    plan,
    result,
    matched_rows,
    insights,
    report,
    anomalies,
    anomaly_explanation,
    selected_report_history
):

    document = Document()

    # -----------------------------------------------------
    # PAGE SETTINGS
    # -----------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.8
    )

    section.right_margin = Inches(
        0.8
    )

    # -----------------------------------------------------
    # DEFAULT FONT
    # -----------------------------------------------------

    styles = document.styles

    styles["Normal"].font.name = (
        "Arial"
    )

    styles["Normal"].font.size = (
        Pt(10)
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title = document.add_heading(
        "AI DATA ANALYST REPORT",
        level=0
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle_run = (
        subtitle.add_run(
            "Automated Business Analysis"
        )
    )

    subtitle_run.bold = True
    subtitle_run.font.size = Pt(12)

    date_paragraph = (
        document.add_paragraph()
    )

    date_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    date_paragraph.add_run(
        datetime.now().strftime(
            "%d %B %Y, %I:%M %p"
        )
    )

    # -----------------------------------------------------
    # EXECUTIVE SUMMARY
    # -----------------------------------------------------

    document.add_heading(
        "1. Executive Summary",
        level=1
    )

    document.add_paragraph(
        str(
            report.get(
                "executive_summary",
                "No executive summary generated."
            )
        )
    )

    # -----------------------------------------------------
    # CURRENT QUESTION
    # -----------------------------------------------------

    document.add_heading(
        "2. Question Analyzed",
        level=1
    )

    question_paragraph = (
        document.add_paragraph()
    )

    question_run = (
        question_paragraph.add_run(
            str(question)
        )
    )

    question_run.bold = True

    # -----------------------------------------------------
    # SELECTED CONVERSATION Q&A
    # -----------------------------------------------------

    if selected_report_history:

        document.add_heading(
            "3. Conversation Questions & Answers",
            level=1
        )

        for index, qa in enumerate(
            selected_report_history,
            start=1
        ):

            document.add_heading(
                f"Question {index}",
                level=2
            )

            document.add_paragraph(
                str(
                    qa.get(
                        "question",
                        ""
                    )
                )
            )

            document.add_heading(
                f"Answer {index}",
                level=2
            )

            document.add_paragraph(
                str(
                    qa.get(
                        "answer",
                        ""
                    )
                )
            )

    # -----------------------------------------------------
    # ANALYSIS DETAILS
    # -----------------------------------------------------

    document.add_heading(
        "4. Analysis Details",
        level=1
    )

    document.add_paragraph(
        "Rows matched: "
        + f"{matched_rows:,}"
    )

    document.add_paragraph(
        "Metric: "
        + str(
            plan.metric
            or "N/A"
        )
    )

    document.add_paragraph(
        "Operation: "
        + str(
            plan.operation
            or "N/A"
        )
    )

    if plan.group_by:

        document.add_paragraph(
            "Grouped by: "
            + str(
                plan.group_by
            )
        )

    if plan.filters:

        document.add_paragraph(
            "Filters applied:"
        )

        for condition in (
            plan.filters
        ):

            document.add_paragraph(
                (
                    f"{condition.column} = "
                    f"{condition.value}"
                ),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "Filters applied: None"
        )

    # -----------------------------------------------------
    # ANALYSIS RESULT
    # -----------------------------------------------------

    document.add_heading(
        "5. Analysis Result",
        level=1
    )

    add_dataframe_to_word(
        document,
        result
    )

    # -----------------------------------------------------
    # CREATE CHART
    # -----------------------------------------------------

    chart_path = None

    if (
        result is not None
        and len(result) > 1
        and len(result.columns) >= 2
    ):

        try:

            x_column = (
                result.columns[0]
            )

            y_column = (
                result.columns[1]
            )

            chart_data = (
                result.copy()
            )

            chart_data[
                y_column
            ] = pd.to_numeric(
                chart_data[
                    y_column
                ],
                errors="coerce"
            )

            chart_data = (
                chart_data.dropna(
                    subset=[
                        y_column
                    ]
                )
            )

            chart_data = (
                chart_data.sort_values(
                    by=y_column,
                    ascending=True
                )
            )

            if len(chart_data) > 15:

                chart_data = (
                    chart_data.tail(15)
                )

            plt.figure(
                figsize=(9, 5)
            )

            plt.barh(
                chart_data[
                    x_column
                ].astype(str),

                chart_data[
                    y_column
                ]
            )

            plt.xlabel(
                str(
                    y_column
                )
                .replace(
                    "_",
                    " "
                )
                .title()
            )

            plt.ylabel(
                str(
                    x_column
                )
                .replace(
                    "_",
                    " "
                )
                .title()
            )

            plt.title(
                "Analysis Result"
            )

            plt.tight_layout()

            chart_path = (
                Path(__file__)
                .resolve()
                .parent
                / "temporary_analyst_chart.png"
            )

            plt.savefig(
                chart_path,
                dpi=200,
                bbox_inches="tight"
            )

            plt.close()

        except Exception:

            chart_path = None

    # -----------------------------------------------------
    # ADD CHART
    # -----------------------------------------------------

    if chart_path is not None:

        document.add_heading(
            "6. Visualization",
            level=1
        )

        document.add_picture(
            str(chart_path),
            width=Inches(6.5)
        )

        last_paragraph = (
            document.paragraphs[-1]
        )

        last_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    # -----------------------------------------------------
    # KEY FINDINGS
    # -----------------------------------------------------

    document.add_heading(
        "7. Key Findings",
        level=1
    )

    findings = report.get(
        "key_findings",
        []
    )

    if findings:

        for finding in findings:

            document.add_paragraph(
                str(finding),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No key findings generated."
        )

    # -----------------------------------------------------
    # TOP PERFORMERS
    # -----------------------------------------------------

    document.add_heading(
        "8. Top Performers",
        level=1
    )

    top_performers = report.get(
        "top_performers",
        []
    )

    if top_performers:

        for item in top_performers:

            document.add_paragraph(
                str(item),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No specific top performers identified."
        )

    # -----------------------------------------------------
    # ANOMALIES
    # -----------------------------------------------------

    document.add_heading(
        "9. Anomalies",
        level=1
    )

    report_anomalies = report.get(
        "anomalies",
        []
    )

    if report_anomalies:

        for item in report_anomalies:

            document.add_paragraph(
                str(item),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No statistical anomalies detected."
        )

    if anomalies is not None:

        add_dataframe_to_word(
            document,
            anomalies,
            "Detected Statistical Anomalies"
        )

        if anomaly_explanation:

            document.add_paragraph(
                "Anomaly Explanation"
            )

            document.add_paragraph(
                str(
                    anomaly_explanation
                )
            )

    # -----------------------------------------------------
    # AI BUSINESS INSIGHTS
    # -----------------------------------------------------

    document.add_heading(
        "10. AI Business Insights",
        level=1
    )

    document.add_paragraph(
        str(insights)
    )

    # -----------------------------------------------------
    # RECOMMENDATIONS
    # -----------------------------------------------------

    document.add_heading(
        "11. Business Recommendations",
        level=1
    )

    recommendations = report.get(
        "recommendations",
        []
    )

    if recommendations:

        for recommendation in (
            recommendations
        ):

            document.add_paragraph(
                str(recommendation),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No recommendations generated."
        )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer

    footer_paragraph = (
        footer.paragraphs[0]
    )

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_paragraph.add_run(
        "Generated by AI Data Analyst"
    )

    # -----------------------------------------------------
    # SAVE REPORT
    # -----------------------------------------------------

    report_path = (
        Path(__file__)
        .resolve()
        .parent
        / "analyst_report.docx"
    )

    document.save(
        report_path
    )

    # -----------------------------------------------------
    # REMOVE TEMP CHART
    # -----------------------------------------------------

    if chart_path is not None:

        try:
            chart_path.unlink()

        except Exception:
            pass

    return report_path


# =========================================================
# 29. FILE UPLOAD
# =========================================================

uploaded_files = st.file_uploader(
    "Upload CSV or Excel files",
    type=[
        "csv",
        "xlsx"
    ],
    accept_multiple_files=True
)

if not uploaded_files:

    st.info(
        "Upload one or more CSV/XLSX files to begin."
    )

    st.stop()


# =========================================================
# 30. READ FILES
# =========================================================

dataframes = []
file_information = []

for uploaded_file in uploaded_files:

    current_df = read_uploaded_file(
        uploaded_file
    )

    if current_df is None:
        continue

    current_df = normalize_dataframe(
        current_df
    )

    dataframes.append(
        current_df
    )

    file_information.append(
        {
            "file": uploaded_file.name,
            "rows": len(current_df),
            "columns": len(
                current_df.columns
            )
        }
    )


if not dataframes:

    st.error(
        "No valid files could be loaded."
    )

    st.stop()


# =========================================================
# 31. SCHEMA VALIDATION
# =========================================================

schema_ok, schema_message = (
    check_schema_compatibility(
        dataframes
    )
)

if not schema_ok:

    st.error(
        "❌ Files have incompatible schemas."
    )

    st.code(
        schema_message
    )

    st.info(
        "Upload files with matching columns "
        "to combine them."
    )

    st.stop()


# =========================================================
# 32. COMBINE DATASETS
# =========================================================

if len(dataframes) == 1:

    df = dataframes[0]

else:

    df = pd.concat(
        dataframes,
        ignore_index=True
    )


# =========================================================
# 33. FILE INFORMATION
# =========================================================

st.subheader(
    "📂 Uploaded Files"
)

file_info_df = pd.DataFrame(
    file_information
)

st.dataframe(
    file_info_df,
    use_container_width=True,
    hide_index=True
)


# =========================================================
# 34. DATASET INFO
# =========================================================

st.subheader(
    "📊 Combined Dataset"
)

col1, col2, col3 = st.columns(3)

with col1:

    st.metric(
        "Files",
        len(dataframes)
    )

with col2:

    st.metric(
        "Total Rows",
        f"{len(df):,}"
    )

with col3:

    st.metric(
        "Columns",
        len(df.columns)
    )


# =========================================================
# 35. DATA PREVIEW
# =========================================================

with st.expander(
    "Preview Combined Dataset"
):

    st.dataframe(
        df.head(20),
        use_container_width=True
    )


# =========================================================
# 36. BASIC STATISTICS
# =========================================================

with st.expander(
    "Basic Statistics"
):

    st.dataframe(
        df.describe(
            include="all"
        ),
        use_container_width=True
    )


# =========================================================
# 37. CONVERSATION CONTROLS
# =========================================================

st.subheader(
    "Conversation"
)

if st.session_state.last_question:

    st.caption(
        "Previous question: "
        + st.session_state.last_question
    )


# =========================================================
# REPORT CHECKBOX
# =========================================================

include_previous_qa = st.checkbox(
    "Include previous questions and answers",
    value=True,
    key="include_previous_qa"
)


# =========================================================
# CLEAR CONVERSATION
# =========================================================

if st.button(
    "🧹 Clear Conversation"
):

    st.session_state.last_question = ""
    st.session_state.last_plan = None
    st.session_state.last_result = None

    st.session_state.chat_history = []

    st.session_state.report_history = []

    st.rerun()


# =========================================================
# 38. QUESTION
# =========================================================

question = st.text_input(
    "Ask a question about your data",
    placeholder=(
        "Example: What is profit by brand?"
    )
)


# =========================================================
# 39. ANALYZE
# =========================================================

if st.button(
    "🔍 Analyze",
    type="primary"
):

    if not question.strip():

        st.warning(
            "Please enter a question."
        )

        st.stop()

    previous_question = (
        st.session_state.last_question
    )

    previous_plan = (
        st.session_state.last_plan
    )

    previous_result = (
        st.session_state.last_result
    )

    # -----------------------------------------------------
    # GEMINI PLAN
    # -----------------------------------------------------

    with st.spinner(
        "Understanding your question..."
    ):

        plan = ask_gemini(
            question,
            df,
            previous_question,
            previous_plan,
            previous_result
        )

    if plan is None:
        st.stop()

    # -----------------------------------------------------
    # FOLLOW-UP
    # -----------------------------------------------------

    if previous_plan:

        resolved_plan = (
            resolve_follow_up(
                question,
                previous_plan,
                previous_question,
                df
            )
        )

        if resolved_plan:
            plan = resolved_plan

    # -----------------------------------------------------
    # IMPROVE PLAN
    # -----------------------------------------------------

    plan = improve_analysis_plan(
        plan,
        question,
        df
    )

    # -----------------------------------------------------
    # SHOW PLAN
    # -----------------------------------------------------

    with st.expander(
        "🔎 Analysis Plan"
    ):

        st.json(
            plan.model_dump()
        )

    # -----------------------------------------------------
    # RUN ANALYSIS
    # -----------------------------------------------------

    try:

        result, matched_rows = (
            run_analysis_plan(
                plan,
                df
            )
        )

    except Exception as e:

        st.error(
            "Analysis failed: "
            + str(e)
        )

        st.stop()

    # -----------------------------------------------------
    # SAVE FOLLOW-UP MEMORY
    # -----------------------------------------------------

    st.session_state.last_question = (
        question
    )

    st.session_state.last_plan = (
        plan.model_dump()
    )

    st.session_state.last_result = (
        result.head(20)
        .to_string(
            index=False
        )
    )

    # -----------------------------------------------------
    # EXPLAINABLE ANALYSIS
    # -----------------------------------------------------

    st.subheader(
        "🧾 Explainable Analysis"
    )

    st.write(
        f"**Rows matched:** "
        f"{matched_rows:,}"
    )

    if plan.filters:

        st.write(
            "**Filters applied:**"
        )

        for condition in (
            plan.filters
        ):

            st.write(
                f"- {condition.column} = "
                f"{condition.value}"
            )

    else:

        st.write(
            "**Filters applied:** None"
        )

    # -----------------------------------------------------
    # RESULT
    # -----------------------------------------------------

    st.subheader(
        "📊 Analysis Result"
    )

    if (
        len(result) == 1
        and len(result.columns) == 1
    ):

        metric_name = (
            result.columns[0]
        )

        value = result.iloc[
            0,
            0
        ]

        if metric_name in {
            "revenue",
            "profit",
            "loss",
            "revenue_per_sale",
            "profit_per_sale"
        }:

            st.metric(
                metric_name
                .replace(
                    "_",
                    " "
                )
                .title(),

                f"₹{value:,.2f}"
            )

        elif metric_name == (
            "profit_margin"
        ):

            st.metric(
                "Profit Margin",
                f"{value:,.2f}%"
            )

        else:

            st.metric(
                metric_name.title(),
                f"{value:,.2f}"
            )

    else:

        st.dataframe(
            result,
            use_container_width=True
        )

    # -----------------------------------------------------
    # VISUALIZATION
    # -----------------------------------------------------

    if len(result) > 1:

        st.subheader(
            "📈 Visualization"
        )

        create_chart(
            result
        )

    # -----------------------------------------------------
    # ANOMALY DETECTION
    # -----------------------------------------------------

    anomalies = detect_anomalies(
        result
    )

    anomaly_explanation = ""

    if anomalies is not None:

        st.subheader(
            "🚨 Anomaly Detection"
        )

        st.write(
            "Potential statistical anomalies "
            "were detected."
        )

        st.dataframe(
            anomalies,
            use_container_width=True
        )

        with st.spinner(
            "Explaining anomalies..."
        ):

            anomaly_explanation = (
                explain_anomalies(
                    question,
                    anomalies
                )
            )

        st.markdown(
            anomaly_explanation
        )

    # -----------------------------------------------------
    # BUSINESS INSIGHTS
    # -----------------------------------------------------

    st.subheader(
        "💡 AI Business Insights"
    )

    with st.spinner(
        "Generating business insights..."
    ):

        insights = (
            generate_business_insights(
                question,
                result
            )
        )

    st.markdown(
        insights
    )

    # -----------------------------------------------------
    # CURRENT ANALYST REPORT
    # -----------------------------------------------------

    st.subheader(
        "📋 Analyst Report"
    )

    with st.spinner(
        "Preparing analyst report..."
    ):

        report = (
            generate_analyst_report(
                question=question,
                plan=plan,
                result=result,
                matched_rows=matched_rows,
                insights=insights,
                anomalies=anomalies,
                anomaly_explanation=(
                    anomaly_explanation
                )
            )
        )

    # =====================================================
    # CREATE ACTUAL ANSWER FOR THIS QUESTION
    # =====================================================

    answer_parts = []

    executive_summary = (
        report.get(
            "executive_summary",
            ""
        )
    )

    if executive_summary:

        answer_parts.append(
            "Executive Summary:\n"
            + str(
                executive_summary
            )
        )

    findings = report.get(
        "key_findings",
        []
    )

    if findings:

        answer_parts.append(
            "Key Findings:\n"
            + "\n".join(
                "- " + str(item)
                for item in findings
            )
        )

    top_performers = report.get(
        "top_performers",
        []
    )

    if top_performers:

        answer_parts.append(
            "Top Performers:\n"
            + "\n".join(
                "- " + str(item)
                for item in top_performers
            )
        )

    report_anomalies = report.get(
        "anomalies",
        []
    )

    if report_anomalies:

        answer_parts.append(
            "Anomalies:\n"
            + "\n".join(
                "- " + str(item)
                for item in report_anomalies
            )
        )

    recommendations = report.get(
        "recommendations",
        []
    )

    if recommendations:

        answer_parts.append(
            "Recommendations:\n"
            + "\n".join(
                "- " + str(item)
                for item in recommendations
            )
        )

    if not answer_parts:

        answer_parts.append(
            "Calculated Result:\n"
            + result.to_string(
                index=False
            )
        )

    current_answer = (
        "\n\n".join(
            answer_parts
        )
    )

    # =====================================================
    # SAVE THIS QUESTION + ANSWER
    # =====================================================

    st.session_state.report_history.append(
        {
            "question": str(
                question
            ),

            "answer": str(
                current_answer
            )
        }
    )

    # =====================================================
    # THIS IS THE IMPORTANT FIX
    # =====================================================

    if include_previous_qa:

        selected_report_history = (
            st.session_state
            .report_history
            .copy()
        )

    else:

        selected_report_history = (
            st.session_state
            .report_history[-1:]
        )

    # -----------------------------------------------------
    # SHOW SELECTED Q&A IN STREAMLIT
    # -----------------------------------------------------

    if selected_report_history:

        st.markdown(
            "### Conversation Questions & Answers"
        )

        for index, qa in enumerate(
            selected_report_history,
            start=1
        ):

            st.markdown(
                f"**Question {index}:**"
            )

            st.write(
                qa["question"]
            )

            st.markdown(
                f"**Answer {index}:**"
            )

            st.write(
                qa["answer"]
            )

            st.divider()

    # -----------------------------------------------------
    # EXECUTIVE SUMMARY
    # -----------------------------------------------------

    st.markdown(
        "### Executive Summary"
    )

    st.write(
        report.get(
            "executive_summary",
            ""
        )
    )

    # -----------------------------------------------------
    # KEY FINDINGS
    # -----------------------------------------------------

    st.markdown(
        "### Key Findings"
    )

    if findings:

        for finding in findings:

            st.markdown(
                "- "
                + str(
                    finding
                )
            )

    else:

        st.write(
            "No key findings generated."
        )

    # -----------------------------------------------------
    # TOP PERFORMERS
    # -----------------------------------------------------

    st.markdown(
        "### Top Performers"
    )

    if top_performers:

        for item in top_performers:

            st.markdown(
                "- "
                + str(
                    item
                )
            )

    else:

        st.write(
            "No specific top performers identified."
        )

    # -----------------------------------------------------
    # ANOMALIES
    # -----------------------------------------------------

    st.markdown(
        "### Anomalies"
    )

    if report_anomalies:

        for item in report_anomalies:

            st.markdown(
                "- "
                + str(
                    item
                )
            )

    else:

        st.write(
            "No statistical anomalies detected."
        )

    # -----------------------------------------------------
    # RECOMMENDATIONS
    # -----------------------------------------------------

    st.markdown(
        "### Business Recommendations"
    )

    if recommendations:

        for recommendation in (
            recommendations
        ):

            st.markdown(
                "- "
                + str(
                    recommendation
                )
            )

    else:

        st.write(
            "No recommendations generated."
        )

    # =====================================================
    # CREATE WORD REPORT
    # =====================================================

    try:

        report_path = (
            create_word_report(
                question=question,
                plan=plan,
                result=result,
                matched_rows=matched_rows,
                insights=insights,
                report=report,
                anomalies=anomalies,
                anomaly_explanation=(
                    anomaly_explanation
                ),

                selected_report_history=(
                    selected_report_history
                )
            )
        )

        st.success(
            "✅ Word report created successfully: "
            + report_path.name
        )

        st.info(
            "The report was generated using "
            "the current checkbox selection."
        )

    except Exception as e:

        st.error(
            "Could not create Word report: "
            + str(e)
        )

    # =====================================================
    # NORMAL CHAT HISTORY
    # =====================================================

    st.session_state.chat_history.append(
        {
            "role": "user",
            "content": str(
                question
            )
        }
    )

    st.session_state.chat_history.append(
        {
            "role": "assistant",
            "content": str(
                current_answer
            )
        }
    )


# =========================================================
# 40. CONVERSATION HISTORY
# =========================================================

if st.session_state.chat_history:

    with st.expander(
        "💬 Conversation History"
    ):

        for message in (
            st.session_state.chat_history
        ):

            if message["role"] == "user":

                st.markdown(
                    "**You:** "
                    + message["content"]
                )

            else:

                st.markdown(
                    "**AI:**"
                )

                st.write(
                    message["content"]
                )