import json
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from config import client, MODEL_NAME
from data_utils import clean_json_response


# =========================================================
# 25. WORD REPORT HELPERS
# =========================================================

def format_value(
    value,
    column_name=""
):

    try:

        numeric_value = float(
            value
        )

    except Exception:

        return str(value)

    column_name = (
        str(column_name)
        .lower()
    )

    if column_name in {
        "revenue",
        "profit",
        "loss",
        "revenue_per_sale",
        "profit_per_sale"
    }:

        return (
            "₹"
            + f"{numeric_value:,.2f}"
        )

    if column_name == "profit_margin":

        return (
            f"{numeric_value:,.2f}%"
        )

    if column_name in {
        "sales",
        "unsold"
    }:

        return (
            f"{numeric_value:,.0f}"
        )

    return (
        f"{numeric_value:,.2f}"
    )


def add_dataframe_to_word(
    document,
    dataframe,
    title=None
):

    if title:

        document.add_heading(
            title,
            level=2
        )

    if dataframe is None:
        return

    if len(dataframe) == 0:

        document.add_paragraph(
            "No data available."
        )

        return

    table = document.add_table(
        rows=1,
        cols=len(
            dataframe.columns
        )
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    header_cells = (
        table.rows[0].cells
    )

    for index, column in enumerate(
        dataframe.columns
    ):

        header_cells[index].text = (
            str(column)
            .replace(
                "_",
                " "
            )
            .title()
        )

    for _, row in (
        dataframe.iterrows()
    ):

        cells = (
            table.add_row().cells
        )

        for index, column in enumerate(
            dataframe.columns
        ):

            cells[index].text = (
                format_value(
                    row[column],
                    column
                )
            )


# =========================================================
# 26. REPORT HISTORY
# =========================================================

def get_report_history_for_output(
    include_previous_qa
):

    history = (
        st.session_state.report_history
    )

    if not history:
        return []

    if include_previous_qa:

        return list(history)

    return [
        history[-1]
    ]


# =========================================================
# 27. GENERATE ANALYST REPORT
# =========================================================

def generate_analyst_report(
    question,
    plan,
    result,
    matched_rows,
    insights
):

    result_text = (
        result.to_string(
            index=False
        )
    )

    prompt = (
        "You are a senior business analyst.\n\n"
        "Create a clean, concise and professional analyst report "
        "using ONLY the information provided below.\n\n"

        "User question:\n"
        + str(question)

        + "\n\n"

        "Analysis plan:\n"
        + json.dumps(
            plan.model_dump(),
            indent=2
        )

        + "\n\n"

        "Rows matched:\n"
        + str(matched_rows)

        + "\n\n"

        "Actual calculated result:\n"
        + result_text

        + "\n\n"

        "AI business insights:\n"
        + str(insights)

        + "\n\n"

        "Return ONLY valid JSON.\n\n"

        "Required structure:\n"

        "{\n"
        '  "executive_summary": "short professional summary",\n'
        '  "key_findings": ["finding 1", "finding 2"],\n'
        '  "top_performers": ["item 1", "item 2"],\n'
        '  "recommendations": ["recommendation 1", "recommendation 2"]\n'
        "}\n\n"

        "Rules:\n"
        "- Do not invent numbers.\n"
        "- Do not calculate new numbers.\n"
        "- Use only supplied information.\n"
        "- Keep findings concise and business-focused.\n"
        "- Recommendations must be based on the supplied analysis."
    )

    try:

        response = (
            client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt
            )
        )

        cleaned = clean_json_response(
            response.text
        )

        report = json.loads(
            cleaned
        )

        if not isinstance(
            report,
            dict
        ):

            raise ValueError(
                "Invalid report format."
            )

        return report

    except Exception as e:

        return {
            "executive_summary": (
                "The analyst report could not "
                "be generated automatically."
            ),

            "key_findings": [
                "Analysis was completed successfully."
            ],

            "top_performers": [],

            "recommendations": [
                "Review the calculated analysis result."
            ],

            "error": str(e)
        }


# =========================================================
# 28. CREATE WORD REPORT
# =========================================================

def create_word_report(
    question,
    plan,
    result,
    matched_rows,
    insights,
    report,
    selected_report_history
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.8
    )

    section.right_margin = Inches(
        0.8
    )

    styles = document.styles

    styles["Normal"].font.name = (
        "Arial"
    )

    styles["Normal"].font.size = (
        Pt(10)
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title = document.add_heading(
        "AI DATA ANALYST REPORT",
        level=0
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle_run = (
        subtitle.add_run(
            "Automated Business Analysis"
        )
    )

    subtitle_run.bold = True
    subtitle_run.font.size = Pt(12)

    date_paragraph = (
        document.add_paragraph()
    )

    date_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    date_paragraph.add_run(
        datetime.now().strftime(
            "%d %B %Y, %I:%M %p"
        )
    )

    # -----------------------------------------------------
    # EXECUTIVE SUMMARY
    # -----------------------------------------------------

    document.add_heading(
        "1. Executive Summary",
        level=1
    )

    document.add_paragraph(
        str(
            report.get(
                "executive_summary",
                "No executive summary generated."
            )
        )
    )

    # -----------------------------------------------------
    # CURRENT QUESTION
    # -----------------------------------------------------

    document.add_heading(
        "2. Question Analyzed",
        level=1
    )

    question_paragraph = (
        document.add_paragraph()
    )

    question_run = (
        question_paragraph.add_run(
            str(question)
        )
    )

    question_run.bold = True

    # -----------------------------------------------------
    # SELECTED CONVERSATION Q&A
    # -----------------------------------------------------

    if selected_report_history:

        document.add_heading(
            "3. Conversation Questions & Answers",
            level=1
        )

        for index, qa in enumerate(
            selected_report_history,
            start=1
        ):

            document.add_heading(
                f"Question {index}",
                level=2
            )

            document.add_paragraph(
                str(
                    qa.get(
                        "question",
                        ""
                    )
                )
            )

            document.add_heading(
                f"Answer {index}",
                level=2
            )

            document.add_paragraph(
                str(
                    qa.get(
                        "answer",
                        ""
                    )
                )
            )

    # -----------------------------------------------------
    # ANALYSIS DETAILS
    # -----------------------------------------------------

    document.add_heading(
        "4. Analysis Details",
        level=1
    )

    document.add_paragraph(
        "Rows matched: "
        + f"{matched_rows:,}"
    )

    document.add_paragraph(
        "Metric: "
        + str(
            plan.metric
            or "N/A"
        )
    )

    document.add_paragraph(
        "Operation: "
        + str(
            plan.operation
            or "N/A"
        )
    )

    if plan.group_by:

        document.add_paragraph(
            "Grouped by: "
            + str(
                plan.group_by
            )
        )

    if plan.filters:

        document.add_paragraph(
            "Filters applied:"
        )

        for condition in (
            plan.filters
        ):

            document.add_paragraph(
                (
                    f"{condition.column} = "
                    f"{condition.value}"
                ),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "Filters applied: None"
        )

    # -----------------------------------------------------
    # ANALYSIS RESULT
    # -----------------------------------------------------

    document.add_heading(
        "5. Analysis Result",
        level=1
    )

    add_dataframe_to_word(
        document,
        result
    )

    # -----------------------------------------------------
    # CREATE CHART
    # -----------------------------------------------------

    chart_path = None

    if (
        result is not None
        and len(result) > 1
        and len(result.columns) >= 2
    ):

        try:

            x_column = (
                result.columns[0]
            )

            y_column = (
                result.columns[1]
            )

            chart_data = (
                result.copy()
            )

            chart_data[
                y_column
            ] = pd.to_numeric(
                chart_data[
                    y_column
                ],
                errors="coerce"
            )

            chart_data = (
                chart_data.dropna(
                    subset=[
                        y_column
                    ]
                )
            )

            chart_data = (
                chart_data.sort_values(
                    by=y_column,
                    ascending=True
                )
            )

            if len(chart_data) > 15:

                chart_data = (
                    chart_data.tail(15)
                )

            plt.figure(
                figsize=(9, 5)
            )

            plt.barh(
                chart_data[
                    x_column
                ].astype(str),

                chart_data[
                    y_column
                ]
            )

            plt.xlabel(
                str(
                    y_column
                )
                .replace(
                    "_",
                    " "
                )
                .title()
            )

            plt.ylabel(
                str(
                    x_column
                )
                .replace(
                    "_",
                    " "
                )
                .title()
            )

            plt.title(
                "Analysis Result"
            )

            plt.tight_layout()

            chart_path = (
                Path(__file__)
                .resolve()
                .parent
                / "temporary_analyst_chart.png"
            )

            plt.savefig(
                chart_path,
                dpi=200,
                bbox_inches="tight"
            )

            plt.close()

        except Exception:

            chart_path = None

    # -----------------------------------------------------
    # ADD CHART
    # -----------------------------------------------------

    if chart_path is not None:

        document.add_heading(
            "5. Visualization",
            level=1
        )

        document.add_picture(
            str(chart_path),
            width=Inches(6.5)
        )

        last_paragraph = (
            document.paragraphs[-1]
        )

        last_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    # -----------------------------------------------------
    # KEY FINDINGS
    # -----------------------------------------------------

    document.add_heading(
        "6. Key Findings",
        level=1
    )

    findings = report.get(
        "key_findings",
        []
    )

    if findings:

        for finding in findings:

            document.add_paragraph(
                str(finding),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No key findings generated."
        )

    # -----------------------------------------------------
    # TOP PERFORMERS
    # -----------------------------------------------------

    document.add_heading(
        "7. Top Performers",
        level=1
    )

    top_performers = report.get(
        "top_performers",
        []
    )

    if top_performers:

        for item in top_performers:

            document.add_paragraph(
                str(item),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No specific top performers identified."
        )

    # -----------------------------------------------------
    # AI BUSINESS INSIGHTS
    # -----------------------------------------------------

    document.add_heading(
        "8. AI Business Insights",
        level=1
    )

    document.add_paragraph(
        str(insights)
    )

    # -----------------------------------------------------
    # RECOMMENDATIONS
    # -----------------------------------------------------

    document.add_heading(
        "9. Business Recommendations",
        level=1
    )

    recommendations = report.get(
        "recommendations",
        []
    )

    if recommendations:

        for recommendation in (
            recommendations
        ):

            document.add_paragraph(
                str(recommendation),
                style="List Bullet"
            )

    else:

        document.add_paragraph(
            "No recommendations generated."
        )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer

    footer_paragraph = (
        footer.paragraphs[0]
    )

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_paragraph.add_run(
        "Generated by AI Data Analyst"
    )

    # -----------------------------------------------------
    # SAVE REPORT
    # -----------------------------------------------------

    report_path = (
        Path(__file__)
        .resolve()
        .parent
        / "analyst_report.docx"
    )

    document.save(
        report_path
    )

    # -----------------------------------------------------
    # REMOVE TEMP CHART
    # -----------------------------------------------------

    if chart_path is not None:

        try:
            chart_path.unlink()

        except Exception:
            pass

    return report_path